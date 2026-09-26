"""Sheet names and header rows of the vendored XLSX export.

scripts/run_planner_export.py calls build_planner.main(), which assembles
the workbook with these writers. This test builds that same sheet list
with empty data rows so CI does not need the external /workspace/eq-legends
tree or a full catalog rebuild. Header labels are the lock; cell values
that are counts or formulas are not.
"""
from __future__ import annotations

import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend" / "vendor"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import build_planner as bp  # noqa: E402
import build_xlsx as bx  # noqa: E402
from openpyxl import Workbook  # noqa: E402

from golden_support import assert_golden  # noqa: E402


def _row_values(ws, row: int) -> list:
    values = []
    for cell in ws[row]:
        values.append(cell.value)
    while values and values[-1] in (None, ""):
        values.pop()
    return values


def _column_a_labels(ws) -> list:
    labels = []
    for row in range(1, (ws.max_row or 1) + 1):
        value = ws.cell(row, 1).value
        if value not in (None, ""):
            labels.append(value)
    return labels


def _alt_header(ws) -> list | None:
    for row in range(1, (ws.max_row or 1) + 1):
        if ws.cell(row, 1).value == "Rank" and ws.cell(row, 2).value == "Item":
            return _row_values(ws, row)
    return None


def build_shape_workbook():
    """Same sheet sequence as build_planner.main(), without catalog rows."""
    summary = {}
    weapon_counts = {"total": 0, "usable": 0}
    gear_counts = {"tri": 0, "merged": 0}
    planner_classes = list(bx.DEFAULT_TRIO)
    wb = Workbook()
    wb.remove(wb.active)
    bp.write_how_to_use_extended(wb, summary, weapon_counts, 0, planner_classes)
    bx.write_class_selector(wb, summary, weapon_counts, gear_counts, planner_classes)
    bp.write_stat_planner(wb, 0)
    bp.write_lookup_sheet(wb.create_sheet("By priority (lookup)"), [], "visible")
    bp.write_lookup_sheet(wb.create_sheet("Max all (lookup)"), [], "visible")
    bp.write_scores_helper(wb.create_sheet("_scores"), [], [])
    bx.write_weapon_sheet(wb.create_sheet("Weapons (by class)"), [], highlight=True)
    trio_label = "+".join(
        {"Paladin": "PAL", "Monk": "MNK", "Wizard": "WIZ"}.get(c, c[:3].upper())
        for c in planner_classes
    )
    bx.write_gear_rows(wb.create_sheet(f"{trio_label} BiS"[:31]), [])
    for cls in bx.ALL_CLASSES:
        bx.write_gear_rows(wb.create_sheet(cls[:31]), [])
    bx.write_gear_rows(wb.create_sheet("All gear (all classes)"), [])
    bx.write_gear_rows(wb.create_sheet("All gear (planner trio)"), [])
    for cls in bx.ALL_CLASSES:
        bx.write_weapon_sheet(wb.create_sheet(f"{cls} weapons top"[:31]), [], highlight=(cls == "Paladin"))
    return wb


def _shape(wb) -> dict:
    sheets = []
    for ws in wb.worksheets:
        entry = {"name": ws.title, "state": ws.sheet_state}
        if ws.title == "How to use":
            entry["title"] = ws["A1"].value
        elif ws.title == "Class selector":
            entry["title"] = ws["A1"].value
            entry["labels"] = _column_a_labels(ws)
        elif ws.title == "Stat planner":
            entry["title"] = ws["A1"].value
            entry["loadout_headers"] = _row_values(ws, 8)
            entry["alternate_headers"] = _alt_header(ws)
        else:
            entry["headers"] = _row_values(ws, 1)
        sheets.append(entry)
    return {"sheet_count": len(sheets), "sheets": sheets}


class XlsxShapeTests(unittest.TestCase):
    def test_sheet_names_and_headers(self):
        shape = _shape(build_shape_workbook())
        names = [sheet["name"] for sheet in shape["sheets"]]
        self.assertEqual(names[0], "How to use")
        self.assertEqual(names[1], "Class selector")
        self.assertEqual(names[2], "Stat planner")
        self.assertIn("PAL+MNK+WIZ BiS", names)
        self.assertIn("_scores", names)
        self.assertEqual(sum(1 for name in names if name.endswith(" weapons top")), 16)
        hidden = {sheet["name"] for sheet in shape["sheets"] if sheet["state"] == "hidden"}
        self.assertIn("_scores", hidden)
        gear = next(sheet for sheet in shape["sheets"] if sheet["name"] == "Wizard")
        self.assertIn("Item", gear["headers"])
        self.assertIn("Haste +0", gear["headers"])
        self.assertIn("Haste +10", gear["headers"])
        weapons = next(sheet for sheet in shape["sheets"] if sheet["name"] == "Weapons (by class)")
        self.assertEqual(weapons["headers"][0], "Class(es)")
        self.assertIn("Haste +0", weapons["headers"])
        assert_golden(self, "xlsx_sheet_headers", shape)


if __name__ == "__main__":
    unittest.main()
